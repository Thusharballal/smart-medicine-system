"""
Inspect kendra.docx paragraph pattern to understand record structure
"""
from docx import Document

def inspect_paragraph_pattern(docx_path):
    """Inspect the paragraph pattern in kendra.docx"""
    
    print("=" * 80)
    print("INSPECTING KENDRA.DOCX PARAGRAPH PATTERN")
    print("=" * 80)
    
    doc = Document(docx_path)
    
    print(f"\nTotal paragraphs: {len(doc.paragraphs)}")
    
    # Show first 50 paragraphs to identify the pattern
    print("\nFirst 50 paragraphs:")
    print("-" * 80)
    
    for idx in range(min(50, len(doc.paragraphs))):
        para = doc.paragraphs[idx]
        text = para.text.strip()
        print(f"{idx:4d}: '{text}'")
    
    # Analyze pattern
    print("\n" + "=" * 80)
    print("PATTERN ANALYSIS")
    print("=" * 80)
    
    # Check if records repeat every N paragraphs
    print("\nAttempting to detect record structure...")
    
    # Look at first few records
    texts = [p.text.strip() for p in doc.paragraphs]
    
    # Find pattern: look for sequential numbers
    sequential_numbers = []
    for idx, text in enumerate(texts[:100]):
        if text.isdigit():
            sequential_numbers.append((idx, int(text)))
    
    print(f"\nSequential numbers found in first 100 paragraphs:")
    for idx, num in sequential_numbers[:10]:
        print(f"  Position {idx}: {num}")
    
    # Calculate gap between sequential numbers
    if len(sequential_numbers) >= 2:
        gaps = []
        for i in range(1, min(10, len(sequential_numbers))):
            gap = sequential_numbers[i][0] - sequential_numbers[i-1][0]
            gaps.append(gap)
        
        if gaps:
            print(f"\nGaps between sequential numbers: {gaps}")
            if len(set(gaps)) == 1:
                fields_per_record = gaps[0]
                print(f"\n✓ CONSISTENT PATTERN DETECTED: {fields_per_record} fields per record")
                
                # Estimate total records
                total_records = len(texts) // fields_per_record
                print(f"✓ ESTIMATED TOTAL RECORDS: {total_records}")
                
                # Show field structure for first record
                print(f"\nFIRST RECORD STRUCTURE:")
                print("-" * 80)
                for i in range(fields_per_record):
                    print(f"  Field {i}: '{texts[i]}'")
                
                # Show field structure for second record
                print(f"\nSECOND RECORD STRUCTURE:")
                print("-" * 80)
                for i in range(fields_per_record):
                    print(f"  Field {i}: '{texts[fields_per_record + i]}'")
                
                return fields_per_record, total_records
    
    return None, None

if __name__ == "__main__":
    docx_path = r"c:\Users\admin\Desktop\Smart Medicine system\kendra.docx"
    inspect_paragraph_pattern(docx_path)
