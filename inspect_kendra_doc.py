"""
Inspect kendra.docx structure to understand the data format
"""
from docx import Document
import json

def inspect_kendra_document(docx_path):
    """Inspect the structure of kendra.docx"""
    
    print("=" * 80)
    print("INSPECTING KENDRA.DOCX")
    print("=" * 80)
    
    doc = Document(docx_path)
    
    # Check for tables
    print(f"\nNumber of tables: {len(doc.tables)}")
    
    if doc.tables:
        for idx, table in enumerate(doc.tables):
            print(f"\n--- Table {idx + 1} ---")
            print(f"Rows: {len(table.rows)}")
            print(f"Columns: {len(table.columns)}")
            
            # Show header row if exists
            if table.rows:
                header_cells = table.rows[0].cells
                print(f"\nFirst row (likely headers):")
                headers = [cell.text.strip() for cell in header_cells]
                for i, header in enumerate(headers):
                    print(f"  Column {i}: '{header}'")
                
                # Show first 5 data rows as samples
                print(f"\nFirst 5 sample records:")
                for row_idx in range(1, min(6, len(table.rows))):
                    row = table.rows[row_idx]
                    print(f"\n  Record {row_idx}:")
                    for col_idx, cell in enumerate(row.cells):
                        cell_text = cell.text.strip()
                        if cell_text:
                            print(f"    {headers[col_idx] if col_idx < len(headers) else f'Column {col_idx}'}: {cell_text}")
    
    # Check for paragraphs
    print(f"\n\nNumber of paragraphs: {len(doc.paragraphs)}")
    if doc.paragraphs:
        print("\nFirst 10 paragraphs:")
        for idx, para in enumerate(doc.paragraphs[:10]):
            text = para.text.strip()
            if text:
                print(f"  {idx}: {text[:100]}...")
    
    # Count total records if table exists
    if doc.tables:
        total_records = len(doc.tables[0].rows) - 1  # Minus header row
        print(f"\n" + "=" * 80)
        print(f"TOTAL KENDRA RECORDS FOUND: {total_records}")
        print("=" * 80)
    
    return doc

if __name__ == "__main__":
    docx_path = r"c:\Users\admin\Desktop\Smart Medicine system\kendra.docx"
    inspect_kendra_document(docx_path)
